"""Phase 6: Salesforce CSV候補行解決処理（Step 1 / CSV candidate resolver）。

設計根拠: docs/PHASE_6_TRIPLE_EVIDENCE_SCHEMA_DESIGN_REVIEW.md 第4版

このモジュールの責務は「店舗コード・selected_business_dateからCSV候補行を検索し、
一意に決まるか・複数あるか・見つからないかを判定する」ことに限定される。
Vision抽出プロンプト・intro_evidence/voice_evidence・正の字読取り・三者照合・
review UIには一切関与しない（それらは別モジュール・別Stepの責務）。

判定はすべて決定論的なPython処理で行い、LLMは使用しない。

重要な注意（レビュー指摘により明記）:
    このモジュールが返す csv_review_required / csv_review_reasons は
    **CSV結合という工程単体の結果**であり、ページ全体としてreview_required
    にすべきかどうかを表すものではない。CSV結合が csv_match_status=unique_match
    かつ csv_review_required=False になったとしても、その後のwritten_total・
    tally_countとの三者比較（Step 4）で不一致が見つかれば、ページ全体としては
    review_requiredになり得る。この区別を呼び出し側は必ず維持すること。

スコープ外（意図的に扱わない）:
    - store_name から store_code への変換（未確認の「入力マスター」シート経由の解決）。
      これは呼び出し側（将来の evidence normalization モジュール）の責務とし、
      本モジュールは既に解決済みの store_code と
      store_code_resolution_source を受け取るだけに留める。
    - written_total_value・tally_count との三者比較（Step 4の責務）。
"""

from __future__ import annotations

import dataclasses
import datetime
import difflib
import logging
import math
import re
from dataclasses import dataclass, field
from enum import Enum
from pathlib import Path
from typing import Callable, Optional

import openpyxl

logger = logging.getLogger(__name__)

__all__ = [
    "COLUMN_STORE_CODE",
    "COLUMN_INTRO_TOTAL",
    "COLUMN_VOICE_CALLOUT_TOTAL",
    "COLUMN_DATE",
    "CsvCandidateResolverError",
    "CsvSourceNotFoundError",
    "CsvSheetNotFoundError",
    "InvalidBusinessDateError",
    "InvalidResolutionRuleError",
    "CsvMatchStatus",
    "StoreCodeResolutionSource",
    "BusinessMode",
    "CsvCandidateRow",
    "NearMatchCandidate",
    "CsvResolverInput",
    "CsvResolutionResult",
    "RESOLUTION_BASIS_UNVERIFIED_REFERENCE",
    "resolve_csv_candidates",
]

# ============================================================
# CSV列定義（Salesforceエクスポート 'CSV'シート、283列）
# 列番号は1始まりの「業務上の列番号」。Pythonのリストインデックス(0始まり)とは
# _business_col_to_index() で変換する。両者を混同しないこと。
# ============================================================
COLUMN_STORE_CODE = 1
COLUMN_INTRO_TOTAL = 36
COLUMN_VOICE_CALLOUT_TOTAL = 105
COLUMN_DATE = 283

_EXCEL_EPOCH = datetime.date(1899, 12, 30)  # openpyxl/Excelのシリアル値変換の基準日


def _business_col_to_index(business_col_number: int) -> int:
    """1始まりの業務上の列番号を0始まりのPythonリストインデックスへ変換する。"""
    if business_col_number < 1:
        raise ValueError(f"業務上の列番号は1以上である必要があります: {business_col_number}")
    return business_col_number - 1


# ============================================================
# 例外（構造化エラー）
# ファイル/シート単位・入力形式の異常はここで明示的に送出する。
# 個々の行データの欠損（店舗コード空欄・日付判読不能等）は例外にせず、
# その行をスキップして CsvResolutionResult.data_quality_notes に記録する。
# ============================================================
class CsvCandidateResolverError(Exception):
    """このモジュールが送出する例外の基底クラス。"""


class CsvSourceNotFoundError(CsvCandidateResolverError):
    """指定されたCSV/Excelファイルが存在しない。"""


class CsvSheetNotFoundError(CsvCandidateResolverError):
    """指定されたsheet_nameがブック内に存在しない。"""


class InvalidBusinessDateError(CsvCandidateResolverError):
    """selected_business_dateがYYYY-MM-DD形式として解釈できない。"""


class InvalidResolutionRuleError(CsvCandidateResolverError):
    """resolution_ruleが指定されているのに resolution_rule_basis
    （監査可能な根拠の説明）が指定されていない、または空文字列である。"""


# ============================================================
# ステータス・列挙型
# ============================================================
class CsvMatchStatus(str, Enum):
    UNIQUE_MATCH = "unique_match"
    MULTIPLE_UNRESOLVED = "multiple_unresolved"
    NOT_FOUND = "not_found"
    RESOLVED_BY_RULE = "resolved_by_rule"
    REFERENCE_MATCH_UNVERIFIED = "reference_match_unverified"


class StoreCodeResolutionSource(str, Enum):
    """店舗コードの情報源。unique_match/resolved_by_ruleへ昇格してよいのは
    DIRECT_FORM_READのみ（設計書2.4節）。それ以外は候補が1件でも
    REFERENCE_MATCH_UNVERIFIEDに留める。"""

    DIRECT_FORM_READ = "direct_form_read"          # 帳票のAU1K欄を直接読み取った値（新帳票）
    UNVERIFIED_REFERENCE = "unverified_reference"   # 店舗名から未確認の参考マスター経由で解決した値（旧帳票）
    UNKNOWN = "unknown"                              # 情報源が呼び出し側から指定されていない（安全側でUNVERIFIED_REFERENCE相当に扱う）


class BusinessMode(str, Enum):
    STORE = "store"
    EVENT = "event"
    UNKNOWN = "unknown"


# ============================================================
# データ型
# ============================================================
@dataclass(frozen=True)
class CsvCandidateRow:
    source_file: str
    sheet_name: str
    row_number: int        # Excel上の1始まり・ヘッダー行込みの行番号
    data_row_index: int    # ヘッダーを除いたデータ行のみの0始まり連番
    store_code: str
    business_date: Optional[str]   # "YYYY-MM-DD"文字列。日付セルが読めない行は候補に含めない
    business_mode: BusinessMode
    intro_total: Optional[int]
    voice_callout_total: Optional[int]


@dataclass(frozen=True)
class NearMatchCandidate:
    source_file: str
    sheet_name: str
    row_number: int
    data_row_index: int
    candidate_store_code: str
    edit_distance: int
    length_difference: int      # len(input_store_code) - len(candidate_store_code)
    difference_summary: str
    intro_total: Optional[int]
    voice_callout_total: Optional[int]


@dataclass(frozen=True)
class CsvResolverInput:
    source_file: str
    sheet_name: str
    selected_business_date: str    # "YYYY-MM-DD"。OCR値ではなく外部設定の正本（設計書4節）
    store_code: Optional[str] = None
    store_name: Optional[str] = None   # 参考情報として保持するのみ。本モジュールは名前→コード変換を行わない
    store_code_resolution_source: StoreCodeResolutionSource = StoreCodeResolutionSource.UNKNOWN

    # business_mode_hintは現状「参考情報として保持するのみ」で、内部ロジックからは一切参照しない。
    # CSV側に店頭/イベントを区別する列が存在しない以上、この値「だけ」を根拠に
    # resolved_by_ruleへ昇格させることは禁止する（レビュー指摘4点目）。
    # resolved_by_ruleを使うには、必ず resolution_rule と resolution_rule_basis を
    # 明示的に渡すこと（下記参照）。
    business_mode_hint: Optional[BusinessMode] = None

    near_match_edit_distance_threshold: int = 2

    # 候補が複数件のとき、監査可能な確定ルールで一意化したい場合にのみ指定する。
    # resolution_ruleを指定する場合、resolution_rule_basis（そのルールの根拠 - 例えば
    # 参照した具体的なCSV列名、または先方確認済みルールの識別子の説明文）と
    # resolution_rule_id（そのルール自体を一意に指す識別子。監査ログでの追跡用）を
    # 必ず併せて指定すること。根拠なしにresolved_by_ruleへ昇格させることを防ぐための
    # 必須トリオとする。
    #
    # 現時点（2026-07-20）では先方確認済みの機械判別ルールが存在しないため、
    # 実運用でresolution_ruleを渡す呼び出しはまだ存在しない想定である。
    # P23のような複数候補ケースは、ルールが確立するまで常にmultiple_unresolvedとなる。
    resolution_rule: Optional[Callable[[list], Optional["CsvCandidateRow"]]] = None
    resolution_rule_basis: Optional[str] = None
    resolution_rule_id: Optional[str] = None


@dataclass(frozen=True)
class CsvResolutionResult:
    csv_match_status: CsvMatchStatus
    csv_candidate_rows: list
    near_match_candidates: list

    # unique_match / resolved_by_rule の場合のみ値を持つ行単位の結果。
    # multiple_unresolved / not_found / reference_match_unverified では常にNone。
    # near_matchは自動解決の対象ではないため、near_matchのみが存在する場合も常にNone。
    resolved_candidate_row: Optional[CsvCandidateRow]

    resolution_basis: str

    # このモジュール（CSV結合工程）単体としてのreview要否。ページ全体のreview_requiredとは別物。
    # モジュールdocstring末尾の「重要な注意」を参照。
    csv_review_required: bool
    csv_review_reasons: list

    data_quality_notes: list = field(default_factory=list)

    def to_dict(self) -> dict:
        """テスト・将来の下流連携向けの単純なdict変換（Enumは.valueに展開）。"""

        def _convert(obj):
            if isinstance(obj, Enum):
                return obj.value
            if dataclasses.is_dataclass(obj):
                return {k: _convert(v) for k, v in dataclasses.asdict(obj).items()}
            if isinstance(obj, list):
                return [_convert(v) for v in obj]
            if isinstance(obj, dict):
                return {k: _convert(v) for k, v in obj.items()}
            return obj

        return _convert(self)


# resolution_basisの定型文言（設計書第4版で指定された文言をそのまま使用する）
RESOLUTION_BASIS_UNVERIFIED_REFERENCE = "ai_store_name_plus_unverified_reference_master"


# ============================================================
# ヘルパー: 正規化
# ============================================================
def _normalize_store_code(raw: Optional[str]) -> Optional[str]:
    """trim・大文字化・半角スペース/ハイフン除去のみ。数字・店舗固有の英字は変換しない。"""
    if raw is None:
        return None
    s = str(raw).strip().upper()
    s = re.sub(r"[ \t\-]", "", s)
    return s or None


def _parse_business_date(value: str) -> datetime.date:
    try:
        return datetime.date.fromisoformat(str(value).strip())
    except (ValueError, AttributeError) as exc:
        raise InvalidBusinessDateError(
            f"selected_business_dateはYYYY-MM-DD形式である必要があります: {value!r}"
        ) from exc


_DATE_STRING_FORMATS = ("%Y-%m-%d", "%Y/%m/%d", "%Y%m%d")


def _excel_serial_to_date(value) -> Optional[datetime.date]:
    """Excelのシリアル値、または日付/文字列表記の値をdatetime.dateへ正規化する。
    解釈できない場合はNoneを返す（例外にしない。呼び出し元でスキップ扱いにする）。"""
    if value is None:
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        try:
            return _EXCEL_EPOCH + datetime.timedelta(days=int(value))
        except (OverflowError, ValueError):
            return None
    if isinstance(value, str):
        s = value.strip()
        for fmt in _DATE_STRING_FORMATS:
            try:
                return datetime.datetime.strptime(s, fmt).date()
            except ValueError:
                continue
        return None
    return None


# \d はUnicode対応のためデフォルトでは全角数字("１２"等)にもマッチしてしまう。
# 半角ASCII数字のみを「整数形式」として許可するため、[0-9]で明示する。
_INTEGER_STRING_RE = re.compile(r"[+-]?[0-9]+")


def _to_optional_int(
    value,
    *,
    column_name: Optional[str] = None,
    row_number: Optional[int] = None,
    notes: Optional[list] = None,
) -> Optional[int]:
    """CSVセルの値を厳格に整数へ変換する。曖昧な値は一切0へフォールバックしない。

    許可する値:
        - None、空文字/空白のみの文字列 -> None（欠損として扱う。0への変換はしない）
        - int -> そのまま許可
        - float かつ is_integer() が真（例: 2.0） -> int()化して許可
        - 整数形式の文字列（例: "2", "-3"。符号のみ・空文字は不可） -> int()化して許可

    無効として None を返す（かつ notes が渡されていれば記録する）値:
        - bool（True/False）。Pythonの bool は int のサブクラスだが、ここでは明示的に無効とする
        - 小数を表す float（例: 2.5、is_integer()が偽）
        - NaN / Infinity（float("nan").is_integer() は False になるため上と同じ経路でNoneになる）
        - 小数形式の文字列（例: "2.0"。整数形式のみ許可し、小数点を含む文字列は拒否する仕様とする）
        - 数字以外の文字を含む文字列、"nan"/"inf"等の特殊文字列
        - 上記以外の未対応の型

    不正な値が見つかっても 0 へは変換しない。呼び出し元がその欄を「欠損」として
    review 対象にできるよう None のまま返す。notes には列名・行番号・理由のみを記録し、
    セル値そのもの（個人情報・機密情報を含み得る）を大量出力しない
    （デバッグ用に先頭数文字のみ切り詰めて含める）。
    """

    def _record(reason: str) -> None:
        if notes is None:
            return
        col = column_name or "?"
        row = row_number if row_number is not None else "?"
        preview = repr(value)
        if len(preview) > 30:
            preview = preview[:30] + "...(truncated)"
        notes.append(f"行{row} 列'{col}': 数値として解釈できないため欠損扱い（理由: {reason}, 値の一部: {preview}）")

    if value is None:
        return None

    if isinstance(value, bool):
        _record("bool型は数値として扱わない")
        return None

    if isinstance(value, int):
        return value

    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            _record("NaN/Infinityは無効な数値")
            return None
        if value.is_integer():
            return int(value)
        _record("小数は整数として扱えない")
        return None

    if isinstance(value, str):
        s = value.strip()
        if s == "":
            return None
        if _INTEGER_STRING_RE.fullmatch(s):
            return int(s)
        _record("整数形式の文字列として解釈できない（小数形式・非数字混在等）")
        return None

    _record(f"未対応の型: {type(value).__name__}")
    return None


# ============================================================
# ヘルパー: 編集距離・差分説明
# ============================================================
def _levenshtein_distance(a: str, b: str) -> int:
    """標準的な動的計画法によるLevenshtein距離（挿入・削除・置換いずれもコスト1）。"""
    if a == b:
        return 0
    if len(a) == 0:
        return len(b)
    if len(b) == 0:
        return len(a)

    prev_row = list(range(len(b) + 1))
    for i, ca in enumerate(a, start=1):
        cur_row = [i] + [0] * len(b)
        for j, cb in enumerate(b, start=1):
            cost = 0 if ca == cb else 1
            cur_row[j] = min(
                prev_row[j] + 1,          # 削除
                cur_row[j - 1] + 1,       # 挿入
                prev_row[j - 1] + cost,   # 置換 or 一致
            )
        prev_row = cur_row
    return prev_row[-1]


def _describe_difference(input_code: str, candidate_code: str, edit_distance: int) -> str:
    """近似候補の差分を人間向けに説明する（監査・参考情報用。厳密な一意の説明を保証するものではない）。"""
    matcher = difflib.SequenceMatcher(None, input_code, candidate_code, autojunk=False)
    ops = [op for op in matcher.get_opcodes() if op[0] != "equal"]

    if len(ops) == 1:
        tag, i1, i2, j1, j2 = ops[0]
        if tag == "insert":
            # candidate側にinputにはない文字がある = inputから見て桁が欠落している
            inserted = candidate_code[j1:j2]
            if len(inserted) == 1:
                ch = inserted
                consecutive = (j1 > 0 and candidate_code[j1 - 1] == ch) or (
                    j2 < len(candidate_code) and candidate_code[j2] == ch
                )
                if consecutive:
                    return f"連続する{ch}の1桁欠落候補"
                return f"1桁欠落候補（候補側の{j1 + 1}文字目に'{ch}'）"
            return f"{len(inserted)}桁欠落候補"
        if tag == "delete":
            removed = input_code[i1:i2]
            if len(removed) == 1:
                return f"1桁余剰候補（入力側の{i1 + 1}文字目に余分な'{removed}'）"
            return f"{len(removed)}桁余剰候補"
        if tag == "replace":
            a_part = input_code[i1:i2]
            b_part = candidate_code[j1:j2]
            if len(a_part) == 1 and len(b_part) == 1:
                return f"1桁置換候補（{i1 + 1}文字目: '{a_part}'→'{b_part}'）"
            return f"置換候補（'{a_part}'→'{b_part}'）"

    return f"編集距離{edit_distance}の近似候補"


def _find_near_matches(
    input_store_code: str,
    date_filtered_rows: list,
    threshold: int,
) -> list:
    """exact matchが0件のときのみ呼び出す。同一selected_business_date内の全店舗コードとの
    編集距離を計算し、閾値以内のものだけをedit_distance昇順で返す。
    近似候補は resolved_candidate_row には反映されない（自動解決の対象外）。"""
    results = []
    for row in date_filtered_rows:
        if row.store_code == input_store_code:
            continue
        dist = _levenshtein_distance(input_store_code, row.store_code)
        if dist <= threshold:
            results.append(
                NearMatchCandidate(
                    source_file=row.source_file,
                    sheet_name=row.sheet_name,
                    row_number=row.row_number,
                    data_row_index=row.data_row_index,
                    candidate_store_code=row.store_code,
                    edit_distance=dist,
                    length_difference=len(input_store_code) - len(row.store_code),
                    difference_summary=_describe_difference(input_store_code, row.store_code, dist),
                    intro_total=row.intro_total,
                    voice_callout_total=row.voice_callout_total,
                )
            )
    results.sort(key=lambda c: (c.edit_distance, c.row_number))
    return results


# ============================================================
# シート読み込み
# ============================================================
def _load_candidate_rows(path: Path, sheet_name: str) -> tuple:
    """CSVシートの全データ行を読み込み、(candidates, data_quality_notes) を返す。
    行単位の欠損（店舗コード空欄・日付判読不能・列数不足）は例外にせず、
    その行をスキップしてnotesに記録する。"""
    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - ファイル破損等を構造化エラーへ変換する
        raise CsvCandidateResolverError(f"ワークブックを開けませんでした: {path} ({exc})") from exc

    try:
        if sheet_name not in wb.sheetnames:
            raise CsvSheetNotFoundError(
                f"シート'{sheet_name}'が見つかりません: {path} (利用可能なシート: {wb.sheetnames})"
            )
        ws = wb[sheet_name]

        store_idx = _business_col_to_index(COLUMN_STORE_CODE)
        intro_idx = _business_col_to_index(COLUMN_INTRO_TOTAL)
        voice_idx = _business_col_to_index(COLUMN_VOICE_CALLOUT_TOTAL)
        date_idx = _business_col_to_index(COLUMN_DATE)
        max_idx = max(store_idx, intro_idx, voice_idx, date_idx)

        candidates: list = []
        notes: list = []
        data_row_index = 0

        # min_row=2 でヘッダー行(1行目)を除外。row_numberはExcel上の実際の行番号(1始まり)。
        for excel_row_number, row in enumerate(ws.iter_rows(min_row=2), start=2):
            values = [c.value for c in row]

            if len(values) <= max_idx:
                notes.append(
                    f"行{excel_row_number}: 列数不足のためスキップ（必要列={max_idx + 1}, 実列数={len(values)}）"
                )
                data_row_index += 1
                continue

            store_code = _normalize_store_code(values[store_idx])
            row_date = _excel_serial_to_date(values[date_idx])

            if store_code is None:
                notes.append(f"行{excel_row_number}: 店舗コードが空欄/判読不能のためスキップ")
                data_row_index += 1
                continue
            if row_date is None:
                notes.append(f"行{excel_row_number}: 日付が判読不能のためスキップ")
                data_row_index += 1
                continue

            candidates.append(
                CsvCandidateRow(
                    source_file=str(path),
                    sheet_name=sheet_name,
                    row_number=excel_row_number,
                    data_row_index=data_row_index,
                    store_code=store_code,
                    business_date=row_date.isoformat(),
                    # CSV列に店頭/イベントを区別する列が存在しないため常にUNKNOWN（設計書2.1節）
                    business_mode=BusinessMode.UNKNOWN,
                    intro_total=_to_optional_int(
                        values[intro_idx],
                        column_name="紹介総数(col36)",
                        row_number=excel_row_number,
                        notes=notes,
                    ),
                    voice_callout_total=_to_optional_int(
                        values[voice_idx],
                        column_name="お声がけ総数(col105)",
                        row_number=excel_row_number,
                        notes=notes,
                    ),
                )
            )
            data_row_index += 1

        return candidates, notes
    finally:
        wb.close()


def _missing_value_review_reasons(row: CsvCandidateRow) -> list:
    """店舗コードの一意解決とは独立に、CSV証拠(数値)が完全かどうかを判定する。
    intro_total・voice_callout_totalのいずれかがNone(欠損)であれば理由を積む。
    実際の値0は欠損として扱わない(_to_optional_intで既に区別済み)。
    両方揃っている場合のみ空リスト(=CSV証拠として完全)を返す。"""
    reasons = []
    if row.intro_total is None:
        reasons.append("intro_total_missing")
    if row.voice_callout_total is None:
        reasons.append("voice_callout_total_missing")
    return reasons


# ============================================================
# メイン関数
# ============================================================
def resolve_csv_candidates(resolver_input: CsvResolverInput) -> CsvResolutionResult:
    """店舗コード・selected_business_dateからCSV候補行を検索し、
    csv_match_status を判定する。Vision抽出・正の字読取り・三者比較は行わない。

    unique_matchへ判定するための条件（すべて満たす場合のみ）:
        1. selected_business_date（外部設定の正本）に一致する行であること
        2. 店舗コードが正規化後に完全一致すること（近似一致では不可）
        3. 該当する候補行がちょうど1件であること
        4. store_code_resolution_source が DIRECT_FORM_READ であること
           （UNVERIFIED_REFERENCE / UNKNOWN は候補1件でも
             REFERENCE_MATCH_UNVERIFIED に留め、昇格させない）
        5. 候補行に店舗コード・日付の欠損がないこと
           （_load_candidate_rows の時点で欠損行は候補から除外済み）

    Returns:
        CsvResolutionResult。csv_review_required/csv_review_reasonsは
        このCSV結合工程単体の結果であり、ページ全体のreview判定ではない
        （モジュールdocstring冒頭を参照）。

    Raises:
        CsvSourceNotFoundError: source_fileが存在しない。
        CsvSheetNotFoundError: sheet_nameがブックに存在しない。
        InvalidBusinessDateError: selected_business_dateがYYYY-MM-DD形式でない。
        InvalidResolutionRuleError: resolution_ruleが指定されているのに
            resolution_rule_basis/resolution_rule_idが指定されていない、
            またはresolution_ruleがexact_matches外の値を返した。
        CsvCandidateResolverError: ワークブックを開く際のその他の異常。
    """
    if resolver_input.resolution_rule is not None:
        if not (resolver_input.resolution_rule_basis or "").strip():
            raise InvalidResolutionRuleError(
                "resolution_ruleを指定する場合はresolution_rule_basis"
                "（監査可能な根拠。参照した具体的なCSV列名や先方確認済みルールの識別子の説明）"
                "を必ず併せて指定してください。business_mode_hintだけを根拠にresolved_by_ruleへ"
                "昇格させることはできません。"
            )
        if not (resolver_input.resolution_rule_id or "").strip():
            raise InvalidResolutionRuleError(
                "resolution_ruleを指定する場合はresolution_rule_id"
                "（そのルール自体を一意に指す識別子。監査ログでの追跡用）"
                "を必ず併せて指定してください。"
            )

    path = Path(resolver_input.source_file)
    if not path.exists():
        raise CsvSourceNotFoundError(f"CSVソースファイルが見つかりません: {path}")

    target_date = _parse_business_date(resolver_input.selected_business_date)

    all_rows, data_quality_notes = _load_candidate_rows(path, resolver_input.sheet_name)
    date_filtered = [r for r in all_rows if r.business_date == target_date.isoformat()]

    input_store_code = _normalize_store_code(resolver_input.store_code)

    if input_store_code is None:
        result = CsvResolutionResult(
            csv_match_status=CsvMatchStatus.NOT_FOUND,
            csv_candidate_rows=[],
            near_match_candidates=[],
            resolved_candidate_row=None,
            resolution_basis="store_code_not_provided",
            csv_review_required=True,
            csv_review_reasons=["store_code_not_provided"],
            data_quality_notes=data_quality_notes,
        )
        _log_result(resolver_input, path, result)
        return result

    exact_matches = [r for r in date_filtered if r.store_code == input_store_code]
    resolution_source = resolver_input.store_code_resolution_source

    # ---- 候補1件 ----
    if len(exact_matches) == 1:
        row = exact_matches[0]
        if resolution_source == StoreCodeResolutionSource.DIRECT_FORM_READ:
            # 店舗コード自体は一意に決まっても、intro_total/voice_callout_totalが
            # 欠損している場合はCSV証拠として不完全なので、csv_review_requiredを立てる。
            # （店舗コードの一意解決 と 数値の完全性 は別軸の判断として分離する）
            missing_reasons = _missing_value_review_reasons(row)
            result = CsvResolutionResult(
                csv_match_status=CsvMatchStatus.UNIQUE_MATCH,
                csv_candidate_rows=[row],
                near_match_candidates=[],
                resolved_candidate_row=row,
                resolution_basis="候補1件のみ（直接読取店舗コードによる一致）",
                csv_review_required=bool(missing_reasons),
                csv_review_reasons=missing_reasons,
                data_quality_notes=data_quality_notes,
            )
        else:
            # UNVERIFIED_REFERENCE または UNKNOWN: 候補1件でもunique_matchへ昇格させない（設計書2.4節）
            result = CsvResolutionResult(
                csv_match_status=CsvMatchStatus.REFERENCE_MATCH_UNVERIFIED,
                csv_candidate_rows=[row],
                near_match_candidates=[],
                resolved_candidate_row=None,
                resolution_basis=RESOLUTION_BASIS_UNVERIFIED_REFERENCE,
                csv_review_required=True,
                csv_review_reasons=["csv_reference_unverified"],
                data_quality_notes=data_quality_notes,
            )
        _log_result(resolver_input, path, result)
        return result

    # ---- 候補2件以上 ----
    if len(exact_matches) >= 2:
        # 2026-07-20時点で先方確認済みの機械判別ルールは存在しない。resolution_ruleを
        # 渡さない限りこの分岐には入らず、P23のような複数候補ケースは常にmultiple_unresolvedになる。
        if resolver_input.resolution_rule is not None:
            resolved_row = resolver_input.resolution_rule(exact_matches)
            if resolved_row is not None:
                # resolution_ruleはexact_matches内の候補行のみを返せる。候補外の値
                # （型違い、または一致する候補が無い値）を返した場合は監査上の異常として扱う。
                if not isinstance(resolved_row, CsvCandidateRow) or resolved_row not in exact_matches:
                    raise InvalidResolutionRuleError(
                        "resolution_ruleはexact_matches内の候補行のみを返せます。"
                        "候補外の行、またはCsvCandidateRow以外の値が返されました。"
                    )
                missing_reasons = _missing_value_review_reasons(resolved_row)
                result = CsvResolutionResult(
                    csv_match_status=CsvMatchStatus.RESOLVED_BY_RULE,
                    csv_candidate_rows=exact_matches,
                    near_match_candidates=[],
                    resolved_candidate_row=resolved_row,
                    resolution_basis=(
                        f"resolved_by_rule[{resolver_input.resolution_rule_id}]: "
                        f"{resolver_input.resolution_rule_basis}"
                    ),
                    csv_review_required=bool(missing_reasons),
                    csv_review_reasons=missing_reasons,
                    data_quality_notes=data_quality_notes,
                )
                _log_result(resolver_input, path, result)
                return result
            # resolved_row is None: ルールが解決を断念した(0件選択に相当) -> multiple_unresolvedへ

        # ルール未指定、またはルールが解決できなかった場合は先頭行等を推測で選ばない
        result = CsvResolutionResult(
            csv_match_status=CsvMatchStatus.MULTIPLE_UNRESOLVED,
            csv_candidate_rows=exact_matches,
            near_match_candidates=[],
            resolved_candidate_row=None,
            resolution_basis=(
                f"候補{len(exact_matches)}件。business_modeを判別する列がCSVに存在しないため一意化不可"
            ),
            csv_review_required=True,
            csv_review_reasons=["csv_multiple_unresolved"],
            data_quality_notes=data_quality_notes,
        )
        _log_result(resolver_input, path, result)
        return result

    # ---- exact_matches == 0件: not_found + 近似探索 ----
    near_matches = _find_near_matches(
        input_store_code=input_store_code,
        date_filtered_rows=date_filtered,
        threshold=resolver_input.near_match_edit_distance_threshold,
    )
    csv_review_reasons = ["csv_not_found"]
    if near_matches:
        csv_review_reasons.append("near_match_available_not_auto_confirmed")

    result = CsvResolutionResult(
        csv_match_status=CsvMatchStatus.NOT_FOUND,
        csv_candidate_rows=[],
        near_match_candidates=near_matches,
        resolved_candidate_row=None,  # 近似候補のみでは自動解決しない
        resolution_basis="exact matchの候補が0件",
        csv_review_required=True,
        csv_review_reasons=csv_review_reasons,
        data_quality_notes=data_quality_notes,
    )
    _log_result(resolver_input, path, result)
    return result


def _log_result(resolver_input: CsvResolverInput, path: Path, result: CsvResolutionResult) -> None:
    """要約のみをログに出す。行データ全体やPIIになり得る情報は出力しない。"""
    logger.info(
        "csv_candidate_resolver: file=%s sheet=%s date=%s status=%s candidates=%d near_matches=%d csv_review_required=%s",
        path.name,
        resolver_input.sheet_name,
        resolver_input.selected_business_date,
        result.csv_match_status.value,
        len(result.csv_candidate_rows),
        len(result.near_match_candidates),
        result.csv_review_required,
    )
