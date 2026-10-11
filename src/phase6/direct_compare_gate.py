"""FAX × Salesforce 直接比較ゲート（2026-10-11）。

対象3項目: 店舗コード / 紹介総数（CSV col36）/ お声がけ総数（CSV col105）。
安全方針: 誤った自動一致（false AUTO_MATCH）を作らないことを最優先する。
- 店舗コードは固定regexで判定しない。有効店舗コード集合（Excel由来マスター）との完全一致のみ。
- 自動補正は一切しない（O↔0、1↔I、桁欠落、余分な文字を勝手に置換しない）。
- 候補は最大3件まで提示するが、必ず NEEDS_REVIEW。
- 自動一致は「observed かつ confidence=high かつ 正解（CSV）一致」のみ。
"""

from __future__ import annotations

import difflib
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Optional

import openpyxl

AUTO_MATCH = "AUTO_MATCH"
NEEDS_REVIEW = "NEEDS_REVIEW"

_WHITESPACE_RE = re.compile(r"[\s　]+")


def normalize_store_code(raw: Optional[str]) -> Optional[str]:
    """前後空白除去 → 内部の空白（半角/全角）除去 → 大文字化。
    文字の置換（O→0 等）は行わない。"""
    if raw is None:
        return None
    s = _WHITESPACE_RE.sub("", str(raw)).upper()
    return s or None


def load_valid_store_codes(xlsx_paths: Iterable[Path]) -> set[str]:
    """Salesforce 日次 Excel（CSV シートの店舗コード列）と入力マスター（店舗コード列）から、
    有効店舗コード集合を読み込む。ファイルごとの日付は問わない（店舗の同一性のみ見る）。"""
    codes: set[str] = set()
    for path in xlsx_paths:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        try:
            for row in wb["CSV"].iter_rows(min_row=2, values_only=True):
                if row and row[0]:
                    norm = normalize_store_code(row[0])
                    if norm:
                        codes.add(norm)
            for row in wb["入力マスター"].iter_rows(min_row=2, values_only=True):
                if row and len(row) > 2 and row[2]:
                    norm = normalize_store_code(row[2])
                    if norm:
                        codes.add(norm)
        finally:
            wb.close()
    return codes


@dataclass(frozen=True)
class StoreCodeValidation:
    raw_value: Optional[str]
    normalized_value: Optional[str]
    exact_master_match: bool
    candidates: tuple  # 最大3件。候補があっても自動確定はしない
    decision: str       # AUTO_MATCH はここでは返さない（判定は decide_item で行う）
    reasons: tuple = field(default_factory=tuple)


def validate_store_code(raw: Optional[str], valid_codes: set[str]) -> StoreCodeValidation:
    norm = normalize_store_code(raw)
    if norm is None:
        return StoreCodeValidation(raw, None, False, (), NEEDS_REVIEW, ("store_code_empty",))
    if norm in valid_codes:
        return StoreCodeValidation(raw, norm, True, (), NEEDS_REVIEW, ())
    cands = difflib.get_close_matches(norm, sorted(valid_codes), n=3, cutoff=0.6)
    return StoreCodeValidation(
        raw, norm, False, tuple(cands), NEEDS_REVIEW, ("store_code_not_in_master",)
    )


@dataclass(frozen=True)
class ItemDecision:
    item: str
    ai_value: Optional[int]
    csv_value: Optional[int]
    decision: str
    reasons: tuple


def decide_written_total(
    *,
    item: str,
    ai_value: Optional[int],
    ai_status: str,
    confidence: Optional[str],
    csv_value: Optional[int],
    csv_row_count: int = 1,
) -> ItemDecision:
    """記載合計（紹介総数 / お声がけ総数）の自動一致判定。
    AUTO_MATCH は: observed・confidence=high・AI値とCSV値が一致・CSV候補行が1件。"""
    reasons = []
    if ai_status != "observed":
        reasons.append(f"ai_status_{ai_status}")
    if ai_status == "observed" and ai_value is None:
        reasons.append("observed_without_value")
    if confidence != "high":
        reasons.append(f"confidence_{confidence}")
    if csv_row_count != 1:
        reasons.append(f"csv_rows_{csv_row_count}")
    if csv_value is None:
        reasons.append("csv_value_missing")
    if ai_value is not None and csv_value is not None and ai_value != csv_value:
        reasons.append("ai_csv_mismatch")
    decision = AUTO_MATCH if not reasons else NEEDS_REVIEW
    return ItemDecision(item, ai_value, csv_value, decision, tuple(reasons))


def decide_store_code(
    *,
    validation: StoreCodeValidation,
    confidence: Optional[str],
    csv_store_code: Optional[str],
) -> ItemDecision:
    """店舗コードの自動一致判定。AUTO_MATCH は完全一致・confidence=high・CSV店舗と一致のみ。"""
    reasons = list(validation.reasons)
    if not validation.exact_master_match:
        reasons.append("no_exact_master_match")
    if confidence != "high":
        reasons.append(f"confidence_{confidence}")
    if csv_store_code is None:
        reasons.append("csv_store_missing")
    elif validation.normalized_value != csv_store_code:
        reasons.append("ai_csv_mismatch")
    decision = AUTO_MATCH if not reasons else NEEDS_REVIEW
    return ItemDecision(
        "store_code", validation.normalized_value, csv_store_code, decision, tuple(reasons)
    )


def apply_manual_override(
    *, item_name: str, manual_value: Optional[int], csv_value: Optional[int]
) -> ItemDecision:
    """人（または手入力相当）の値を受け付けて再照合する。自動一致ではなく手入力確認の結果。
    注意: 本関数のテストは『手入力値を受け付けて再照合するロジック』の検証であり、
    人による確認が行われたことを意味しない。"""
    if manual_value is None or csv_value is None:
        return ItemDecision(item_name, manual_value, csv_value, NEEDS_REVIEW, ("manual_or_csv_missing",))
    if manual_value == csv_value:
        return ItemDecision(item_name, manual_value, csv_value, "MANUAL_CONFIRMED_MATCH", ())
    return ItemDecision(item_name, manual_value, csv_value, NEEDS_REVIEW, ("manual_csv_mismatch",))
